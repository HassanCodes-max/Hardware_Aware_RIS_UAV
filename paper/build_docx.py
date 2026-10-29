"""Build an IEEE conference-style Word file from paper/revised_paper.md.

Layout follows the IEEE conference template: US Letter, Times New Roman,
single-column title block, two-column body, table captions above tables,
figure captions below figures. Wide tables and Fig. 4 span both columns.
[VERIFY] items are highlighted so they are not missed before submission.

    python paper/build_docx.py   ->   paper/revised_paper_ieee.docx
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image

PAPER_DIR = Path(__file__).resolve().parent
SOURCE = PAPER_DIR / "revised_paper.md"
OUTPUT = PAPER_DIR / "revised_paper_ieee.docx"

FONT = "Times New Roman"
COLUMN_WIDTH_IN = 3.5
FULL_WIDTH_IN = 7.25
# Figures drawn wider than this (inches, from the PNG's dpi) span both columns.
WIDE_FIGURE_MIN_IN = 1.5 * COLUMN_WIDTH_IN
# Average 8 pt Times character width, used to decide if a table fits a column.
TABLE_CHAR_WIDTH_IN = 0.055
TABLE_CELL_PADDING_IN = 0.07
HEADER_BOLD_FACTOR = 1.15
NBSP = "\u00a0"

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
SCRIPT = re.compile(r"_\{([^}]*)\}|\^\(([^()]*)\)|_(\w+)|\^(\w+)")
EQUATION = re.compile(r"^(.+?)\s{2,}\((\d+)\)$")
CAPTION = re.compile(r"^\*((Table|Fig\.) [IVX0-9]+)\.\s+(.*)\*$", re.S)
IMAGE = re.compile(r"^!\[[^\]]*\]\(([^)]+)\)$")
LIST_ITEM = re.compile(r"^(\d+)\.\s+(.*)$")


def add_text(paragraph, text, size, bold=False, italic=False, small_caps=False):
    """Add text with _sub / ^sup markup as runs."""

    position = 0
    for match in SCRIPT.finditer(text):
        if match.start() > position:
            _run(paragraph, text[position:match.start()], size, bold, italic, small_caps)
        sub = match.group(1) or match.group(3)
        sup = match.group(2) or match.group(4)
        run = _run(paragraph, sub or sup, size, bold, italic, small_caps)
        if sub:
            run.font.subscript = True
        else:
            run.font.superscript = True
        position = match.end()
    if position < len(text):
        _run(paragraph, text[position:], size, bold, italic, small_caps)


def _run(paragraph, text, size, bold, italic, small_caps):
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.small_caps = small_caps
    if "VERIFY" in text:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def add_inline(paragraph, text, size, bold=False, italic=False, small_caps=False):
    """Add text with **bold**, *italic* and `code` markdown."""

    for token in INLINE.split(text):
        if not token:
            continue
        if token.startswith("**"):
            add_text(paragraph, token[2:-2], size, True, italic, small_caps)
        elif token.startswith("*"):
            add_text(paragraph, token[1:-1], size, bold, not italic, small_caps)
        elif token.startswith("`"):
            _run(paragraph, token[1:-1], size, bold, italic, small_caps).font.name = "Consolas"
        else:
            add_text(paragraph, token, size, bold, italic, small_caps)


def paragraph_format(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0,
                     first_line=None, left=None, hanging=None, keep_next=False):
    fmt = paragraph.paragraph_format
    paragraph.alignment = align
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if first_line is not None:
        fmt.first_line_indent = Inches(first_line)
    if left is not None:
        fmt.left_indent = Inches(left)
    if hanging is not None:
        fmt.left_indent = Inches(hanging)
        fmt.first_line_indent = Inches(-hanging)
    fmt.keep_with_next = keep_next
    return paragraph


def set_columns(section, count):
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        section._sectPr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(int(0.25 * 1440)))


class Builder:
    def __init__(self):
        self.doc = Document()
        self.columns = 1
        self._setup()

    def _setup(self):
        section = self.doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin, section.bottom_margin = Inches(0.75), Inches(1.0)
        section.left_margin = section.right_margin = Inches(0.625)
        set_columns(section, 1)

        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(10)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.line_spacing = 1.0

    def switch_columns(self, count):
        if count == self.columns:
            return
        section = self.doc.add_section(WD_SECTION.CONTINUOUS)
        set_columns(section, count)
        self.columns = count

    @property
    def width_in(self):
        return COLUMN_WIDTH_IN if self.columns == 2 else FULL_WIDTH_IN

    def title(self, text):
        p = paragraph_format(self.doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER, after=12)
        add_inline(p, text, 24)

    def author(self, text):
        name, _, affiliation = text.partition(",")
        p = paragraph_format(self.doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
        add_inline(p, name.replace("**", ""), 11)
        parts = [part.strip() for part in affiliation.split(",")]
        lines = parts[:-2] + [", ".join(parts[-2:])]
        for i, line in enumerate(lines):
            p = paragraph_format(self.doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER,
                                 after=18 if i == len(lines) - 1 else 0)
            add_inline(p, line, 10, italic=True)

    def abstract(self, label, text):
        p = paragraph_format(self.doc.add_paragraph(), first_line=0.2, after=6)
        add_inline(p, label + "—", 9, bold=True, italic=True)
        add_inline(p, text, 9, bold=True)

    def heading(self, text, level):
        if level == 2:
            p = paragraph_format(self.doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER,
                                 before=8, after=4, keep_next=True)
            add_inline(p, text, 10, small_caps=True)
        else:
            p = paragraph_format(self.doc.add_paragraph(), WD_ALIGN_PARAGRAPH.LEFT,
                                 before=4, after=2, keep_next=True)
            add_inline(p, text, 10, italic=True)

    def body(self, text, size=10):
        p = paragraph_format(self.doc.add_paragraph(), first_line=0.2)
        add_inline(p, text, size)

    def list_item(self, number, text):
        p = paragraph_format(self.doc.add_paragraph(), hanging=0.25)
        p.paragraph_format.left_indent = Inches(0.45)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.45))
        add_inline(p, f"{number})\t{text}", 10)

    def reference(self, text):
        p = paragraph_format(self.doc.add_paragraph(), hanging=0.3)
        label, _, rest = text.partition("] ")
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
        add_inline(p, f"{label}]\t{rest}", 8)

    def equation(self, parts, number):
        width = self.width_in
        for i, part in enumerate(parts):
            p = paragraph_format(self.doc.add_paragraph(), WD_ALIGN_PARAGRAPH.LEFT,
                                 before=4 if i == 0 else 0,
                                 after=4 if i == len(parts) - 1 else 0)
            stops = p.paragraph_format.tab_stops
            stops.add_tab_stop(Inches(width / 2), WD_TAB_ALIGNMENT.CENTER)
            stops.add_tab_stop(Inches(width), WD_TAB_ALIGNMENT.RIGHT)
            add_text(p, "\t" + part, 10, italic=False)
            if i == 0:
                add_text(p, f"\t({number})", 10)

    def figure(self, path, caption_label, caption_text):
        image = PAPER_DIR / path
        with Image.open(image) as img:
            drawn_width_in = img.width / img.info.get("dpi", (300, 300))[0]
        wide = drawn_width_in > WIDE_FIGURE_MIN_IN
        if wide:
            self.switch_columns(1)
        width = (self.width_in - 0.1) if not wide else 6.5
        # A one-cell table whose row cannot split keeps the image and its
        # caption in the same column.
        frame = self.doc.add_table(rows=1, cols=1)
        frame.alignment = WD_TABLE_ALIGNMENT.CENTER
        _cant_split(frame.rows[0])
        cell = frame.cell(0, 0)
        p = paragraph_format(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER, before=4)
        p.add_run().add_picture(str(image), width=Inches(width))
        p = paragraph_format(cell.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER, after=6)
        add_inline(p, caption_label + ". ", 8)
        add_inline(p, caption_text, 8)
        paragraph_format(self.doc.add_paragraph(), after=2)
        if wide:
            self.switch_columns(2)

    def table(self, label, caption, rows):
        header, body = rows[0], rows[1:]
        # Approximate width each column needs: its longest body cell, or the
        # longest header word since headers may wrap between words.
        # Keep units in parentheses on one line, e.g. "(16 cm)".
        rows = [[re.sub(r"\(([^)]*)\)", lambda m: "(" + m.group(1).replace(" ", NBSP) + ")", cell)
                 for cell in row] for row in rows]
        header, body = rows[0], rows[1:]
        needs = [
            max([len(row[c].replace("*", "")) for row in body]
                + [HEADER_BOLD_FACTOR * len(word) for word in header[c].replace("*", "").split(" ")])
            * TABLE_CHAR_WIDTH_IN + TABLE_CELL_PADDING_IN
            for c in range(len(header))
        ]
        wide = sum(needs) > COLUMN_WIDTH_IN
        # Use the spare width of a full-width table to keep headers on one line.
        unwrapped = [
            max(need, HEADER_BOLD_FACTOR * len(header[c].replace("*", "")) * TABLE_CHAR_WIDTH_IN
                + TABLE_CELL_PADDING_IN)
            for c, need in enumerate(needs)
        ]
        if sum(unwrapped) <= (FULL_WIDTH_IN if wide else COLUMN_WIDTH_IN) - 0.25:
            needs = unwrapped
        if wide:
            self.switch_columns(1)
        width = self.width_in - (0.05 if not wide else 0.25)

        p = paragraph_format(self.doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER,
                             before=6, keep_next=True)
        add_inline(p, label.upper(), 8)
        p = paragraph_format(self.doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER,
                             after=3, keep_next=True)
        add_inline(p, caption, 8, small_caps=True)

        widths = [width * need / sum(needs) for need in needs]

        table = self.doc.add_table(rows=len(rows), cols=len(header))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _table_borders(table)
        text_columns = [
            sum(bool(re.match(r"\W*[A-Za-z]", row[c])) for row in body) > len(body) / 2
            for c in range(len(header))
        ]
        for r, row in enumerate(rows):
            for c, cell_text in enumerate(row):
                cell = table.cell(r, c)
                cell.width = Inches(widths[c])
                p = cell.paragraphs[0]
                left = c == 0 or text_columns[c]
                paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER)
                add_inline(p, cell_text, 8, bold=(r == 0))
                if r == 0:
                    _cell_bottom_border(cell)
        _set_grid(table, widths)
        p = paragraph_format(self.doc.add_paragraph(), after=6)
        if wide:
            self.switch_columns(2)

    def save(self, path):
        self.doc.save(path)


def _table_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    table._tbl.tblPr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), "50")
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    table._tbl.tblPr.append(margins)


def _cant_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _cell_bottom_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    tc_pr.append(borders)


def _set_grid(table, widths):
    grid = table._tbl.tblGrid
    for col, width in zip(grid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(width * 1440)))


def parse_blocks(text):
    """Split markdown into blocks of joined lines, stopping at the reproduction map."""

    text = text.split("## Reproducing this paper")[0]
    blocks, current = [], []
    for line in text.splitlines():
        if not line.strip():
            if current:
                blocks.append(current)
            current = []
        else:
            current.append(line)
    if current:
        blocks.append(current)
    return blocks


def table_rows(lines):
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", cell) for cell in cells):
            continue
        rows.append(cells)
    return rows


def build(source=SOURCE, output=OUTPUT):
    builder = Builder()
    section = None
    pending_caption = None
    pending_image = None
    seen_author = False

    for lines in parse_blocks(source.read_text(encoding="utf-8")):
        first = lines[0]
        joined = " ".join(line.strip() if i else line for i, line in enumerate(lines))

        if first == "---":
            continue
        if first.startswith("# "):
            builder.title(first[2:])
            continue
        if first.startswith("*Revised draft"):
            continue
        if not seen_author and first.startswith("**"):
            builder.author(joined)
            seen_author = True
            continue
        if first.startswith("## "):
            section = first[3:]
            if section == "Abstract":
                builder.switch_columns(2)
                continue
            builder.heading(section, 2)
            continue
        if first.startswith("### "):
            builder.heading(first[4:], 3)
            continue

        if section == "Abstract":
            if joined.startswith("**Index Terms**"):
                builder.abstract("Index Terms", joined.split(":", 1)[1].strip())
            else:
                builder.abstract("Abstract", joined)
            continue

        if first.startswith("|"):
            label, caption = pending_caption
            builder.table(label, caption, table_rows(lines))
            pending_caption = None
            continue

        image = IMAGE.match(joined)
        if image:
            pending_image = image.group(1)
            continue

        caption = CAPTION.match(joined)
        if caption:
            label, kind, text = caption.group(1), caption.group(2), caption.group(3)
            if kind == "Fig.":
                builder.figure(pending_image, label, text)
                pending_image = None
            else:
                pending_caption = (label, text.rstrip("."))
            continue

        if section == "References":
            builder.reference(joined)
            continue

        equation = EQUATION.match(joined)
        if equation:
            parts = re.split(r"\s{3,}", equation.group(1))
            builder.equation(parts, equation.group(2))
            continue

        if LIST_ITEM.match(first):
            items, current = [], None
            for line in lines:
                item = LIST_ITEM.match(line)
                if item:
                    current = [item.group(1), item.group(2)]
                    items.append(current)
                else:
                    current[1] += " " + line.strip()
            for number, text in items:
                builder.list_item(number, text)
            continue

        builder.body(joined)

    # A closing continuous break balances the last page's two columns.
    builder.switch_columns(1)
    builder.save(output)
    return output


if __name__ == "__main__":
    print(build())
